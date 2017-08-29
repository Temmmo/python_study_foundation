#! python3
import docx,os,openpyxl
os.getcwd()
doc = docx.Document('anqiao.docx')

table = openpyxl.load_workbook('test2.xlsx')

sheet = table.get_sheet_by_name('Sheet1')

Text=[]

for para in doc.paragraphs:
    Text.append(para.text)
for i in range(start, len(Text)+start):
    sheet['A' + str(i)] = Text[0]



